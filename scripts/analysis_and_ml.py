import pandas as pd, numpy as np, openpyxl, matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D = "../data/"; O = "../output/"
GOLD, SAGE, GREY = "#c99a3b", "#6fa88f", "#8a93a6"

# ---------- Excel: guest relations tour tracker ----------
df = pd.read_csv(D + "tour_log.csv")
wb = openpyxl.Workbook(); rd = wb.active; rd.title = "README"
tf, bf, nf = Font(name="Arial", size=14, bold=True), Font(name="Arial", size=11, bold=True), Font(name="Arial", size=11)
lines = [("Guest Relations Tour & Delegation Tracker (simulated data)", tf), ("", nf),
 (f"{len(df)} records, Apr 2019 - Oct 2024: campus tours, admission counselling, delegations and the", nf),
 ("World Congress of Poets (Sep 2019) across KIIT, KIMS and KISS.", nf),
 ("Rows marked Real are the May 2019 admission counselling days reconstructed from memory.", nf),
 ("Everything else, including all names and event numbers, is simulated for practice.", nf), ("", nf),
 ("Pivot exercises", bf),
 ("1. Rows: Visitor_Type | Values: Sum of Group_Size", nf),
 ("2. Rows: Tour_Date (group by Month) | Values: Count", nf),
 ("3. Rows: Campuses_Covered | Values: Average of Satisfaction_Score", nf),
 ("4. Rows: Visitor_Type, Columns: Outcome | Values: Count", nf),
 ("5. Rows: Visitor_Country | Values: Count", nf)]
for i, (t, f) in enumerate(lines, 1): rd.cell(row=i, column=1, value=t).font = f
rd.column_dimensions["A"].width = 95
ws = wb.create_sheet("Tour_Log"); hs = list(df.columns)
hf = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
th = Side(style="thin", color="D9D9D9"); bd = Border(left=th, right=th, top=th, bottom=th)
for c, h in enumerate(hs, 1):
    x = ws.cell(row=1, column=c, value=h); x.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    x.fill = hf; x.alignment = Alignment(horizontal="center", wrap_text=True); x.border = bd
ws.freeze_panes = "A2"; si = hs.index("Data_Source")
for r, row in enumerate(df.itertuples(index=False), 2):
    for c, v in enumerate(row, 1):
        x = ws.cell(row=r, column=c, value=v); x.font = Font(name="Arial", size=10); x.border = bd
    ws.cell(row=r, column=si + 1).fill = PatternFill(start_color="DCEEE3" if row[si] == "Real" else "FBEFD5",
                                                    end_color="DCEEE3" if row[si] == "Real" else "FBEFD5", fill_type="solid")
for i, h in enumerate(hs, 1): ws.column_dimensions[get_column_letter(i)].width = {"Visitor_Type":34,"Campuses_Covered":22,"Purpose":36,"Outcome":26,"Lead_Visitor_Name":22,"Visitor_Title":26}.get(h, 14)
wb.save(O + "01_Excel_Guest_Relations_Tour_Tracker.xlsx")

# ---------- Python EDA: IROP + baggage charts ----------
irop = pd.read_csv(D + "irop_events.csv", parse_dates=["Event_Date"])
bag = pd.read_csv(D + "baggage_cases.csv", parse_dates=["Report_Date"])
kims = pd.read_csv(D + "kims_covid_occupancy.csv", parse_dates=["Date"])
kg = pd.read_csv(D + "kiss_growth.csv")
def style(ax): ax.spines[["top", "right"]].set_visible(False)

m = irop.groupby(irop["Event_Date"].dt.month).size()
fig, ax = plt.subplots(figsize=(8, 4)); ax.bar(m.index, m.values, color=[GOLD if k in (12, 1) else SAGE if k in (6,7,8,9) else GREY for k in m.index])
ax.set_xticks(range(1, 13)); ax.set_title("Disruption events by month (gold = fog, green = monsoon)", fontsize=11)
ax.set_xlabel("Month"); ax.set_ylabel("Events"); style(ax); plt.tight_layout(); plt.savefig(O + "03_chart_irop_by_month.png", dpi=150); plt.close()

w = irop.groupby("Disruption_Type")["Avg_Wait_Minutes"].mean().sort_values()
fig, ax = plt.subplots(figsize=(6, 3.8)); w.plot(kind="barh", ax=ax, color=GOLD)
ax.set_title("Average passenger wait by disruption type", fontsize=11); ax.set_xlabel("Minutes"); style(ax)
plt.tight_layout(); plt.savefig(O + "03_chart_irop_wait_by_type.png", dpi=150); plt.close()

ot = bag.groupby("Issue_Type")["Closed_On_Time"].mean().mul(100)
fig, ax = plt.subplots(figsize=(6, 3.8)); ot.plot(kind="bar", ax=ax, color=SAGE, rot=0)
ax.set_title("Baggage cases closed on time (%)", fontsize=11); ax.set_ylabel("%"); style(ax)
plt.tight_layout(); plt.savefig(O + "03_chart_baggage_on_time.png", dpi=150); plt.close()

fig, ax = plt.subplots(figsize=(9, 4)); ax.plot(kims["Date"], kims["Occupancy_Rate"] * 100, color=GOLD)
ax.axhline(85, color=GREY, ls="--", lw=1, label="85% warning line"); ax.legend(fontsize=8)
ax.set_title("Hospital bed occupancy, Oct 2020 - Oct 2021 (simulated)", fontsize=11); ax.set_ylabel("%"); style(ax)
plt.tight_layout(); plt.savefig(O + "04_chart_kims_occupancy.png", dpi=150); plt.close()

fig, ax = plt.subplots(figsize=(8, 4.2)); ax.plot(kg["Year"], kg["Students"], marker="o", color=GOLD)
for _, r in kg.iterrows(): ax.annotate(f"{r.Students:,}", (r.Year, r.Students), textcoords="offset points", xytext=(0, 9), ha="center", fontsize=8)
ax.set_title("KISS growth: 125 students (1992) to 30,000+", fontsize=11); style(ax)
plt.tight_layout(); plt.savefig(O + "04_chart_kiss_growth.png", dpi=150); plt.close()

for src, dst in [("tour_log", "01_PowerBI_tour_log"), ("irop_events", "03_PowerBI_irop_events"),
                 ("baggage_cases", "03_PowerBI_baggage_cases"), ("kims_covid_occupancy", "04_PowerBI_kims_occupancy"),
                 ("kiss_growth", "04_PowerBI_kiss_growth")]:
    pd.read_csv(D + src + ".csv").to_csv(O + dst + ".csv", index=False)

# ---------- ML: predict late baggage case closure ----------
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import OneHotEncoder
from sklearn.compose import ColumnTransformer
from sklearn.pipeline import Pipeline
from sklearn.linear_model import LogisticRegression
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score
bag["Report_Month"] = bag["Report_Date"].dt.month
bag["Late"] = 1 - bag["Closed_On_Time"]
X = bag[["Issue_Type", "Connecting_Passenger", "Found_At_First_Station", "Peak_Month", "Report_Month"]]; y = bag["Late"]
Xtr, Xte, ytr, yte = train_test_split(X, y, test_size=0.25, random_state=42, stratify=y)
prep = ColumnTransformer([("c", OneHotEncoder(handle_unknown="ignore"), ["Issue_Type"])], remainder="passthrough")
res = {}
for name, mdl in [("Logistic Regression", LogisticRegression(max_iter=1000)), ("Random Forest", RandomForestClassifier(n_estimators=200, random_state=42))]:
    p = Pipeline([("p", prep), ("m", mdl)]).fit(Xtr, ytr); pr = p.predict(Xte)
    res[name] = [accuracy_score(yte, pr), precision_score(yte, pr), recall_score(yte, pr), f1_score(yte, pr)]
base = max(yte.mean(), 1 - yte.mean())
with open(O + "05_ML_model_card.md", "w") as f:
    f.write("# Late Baggage Case Predictor: Model Card\n\n")
    f.write("**Task:** predict whether a mishandled-baggage case will be closed late (past its target days) from issue type, connecting passenger, whether the bag was found at the first station, and season.\n\n")
    f.write(f"**Data:** {len(bag)} simulated cases (Oct 2014 - Mar 2019); {y.mean()*100:.1f}% closed late. Simulated, so results show the method, not real airline performance.\n\n")
    f.write(f"**Baseline (always predict the majority class):** {base:.3f} accuracy\n\n")
    f.write("| Model | Accuracy | Precision | Recall | F1 |\n|---|---|---|---|---|\n")
    for k, v in res.items(): f.write(f"| {k} | {v[0]:.3f} | {v[1]:.3f} | {v[2]:.3f} | {v[3]:.3f} |\n")
    f.write("\n**Reading it:** compare accuracy to the baseline, not to 100%. Recall matters most here, because the point is to catch cases likely to run late so someone can chase them early.\n")
print({k: [round(x, 3) for x in v] for k, v in res.items()}, "baseline", round(base, 3))
